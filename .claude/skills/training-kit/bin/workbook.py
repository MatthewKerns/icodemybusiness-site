#!/usr/bin/env python3
"""Workbook builder — JSON spec → house-style .xlsx, uploaded by the Drive connector as a native Google Sheet.

  workbook.py <spec.json> [-o out.xlsx]     build the .xlsx and print the upload call. Never uploads.
  workbook.py --example                     print a starter spec

Spec:
  {"title": "1.4 Time slicing — tracker",
   "sheets": [
     {"name": "Log", "instructions": "One row per block. Fill the gold columns.",
      "columns": [{"header": "Date", "width": 12, "type": "date"},
                  {"header": "Objective", "width": 28, "input": true, "choices": ["Client work", "Marketing", "Admin"]},
                  {"header": "Minutes", "width": 10, "input": true, "type": "int"},
                  {"header": "Hours", "width": 10, "formula": "=C{r}/60", "format": "0.00"}],
      "rows": 30,                      # blank input rows to pre-format (formulas filled down)
      "data": [["2026-10-02", "Marketing", 45]],   # optional pre-filled rows (cited examples only)
      "totals": {"Minutes": "SUM", "Hours": "SUM"}}],
   "sources": ["[BT] Pillar1-Clockify-Basic-2026-09-13.txt", "[CLK-004]"]}   # → a Sources sheet

House style: black header with gold text, frozen header row, gold-tinted input columns, dropdowns
for `choices`, filled-down formulas, a totals row, and a Sources sheet so every number traces back.
Upload (SKILL.md § Workbooks): connector create_file {title, base64Content=<xlsx b64>,
contentMimeType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", parentId=<folder>}
— Drive converts it to a Google Sheet. Members get the /copy link; sharing changes are Matthew's gate.
"""
import datetime, json, os, sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BLACK, GOLD, INPUT, LINE = "0A0A0A", "D4AF37", "FBF5E1", "D9D4C7"
EXAMPLE = {
    "title": "1.4 Time slicing — tracker",
    "sheets": [{"name": "Log", "instructions": "One row per block. Fill the gold columns.",
                "columns": [{"header": "Date", "width": 12, "type": "date"},
                            {"header": "Objective", "width": 28, "input": True, "choices": ["Client work", "Marketing", "Admin"]},
                            {"header": "Minutes", "width": 10, "input": True, "type": "int"},
                            {"header": "Hours", "width": 10, "formula": "=C{r}/60", "format": "0.00"}],
                "rows": 30, "data": [["2026-10-02", "Marketing", 45]],
                "totals": {"Minutes": "SUM", "Hours": "SUM"}}],
    "sources": ["[BT] Pillar1-Clockify-Basic-2026-09-13.txt", "[CLK-004]"],
}


def build(spec, out):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color=LINE)
    for sh in spec["sheets"]:
        ws = wb.create_sheet(sh["name"][:31])
        cols, top = sh["columns"], 1
        if sh.get("instructions"):
            ws.cell(1, 1, sh["instructions"]).font = Font(italic=True, color="6B665C")
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
            top = 2
        for c, col in enumerate(cols, 1):
            cell = ws.cell(top, c, col["header"])
            cell.font, cell.fill = Font(bold=True, color=GOLD), PatternFill("solid", fgColor=BLACK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = col.get("width", 16)
        ws.freeze_panes = ws.cell(top + 1, 1)
        data, n = sh.get("data", []), max(sh.get("rows", 0), len(sh.get("data", [])))
        for c, col in enumerate(cols, 1):
            if col.get("choices"):
                dv = DataValidation(type="list", formula1='"' + ",".join(col["choices"]) + '"', allow_blank=True)
                ws.add_data_validation(dv)
                dv.add(f"{get_column_letter(c)}{top + 1}:{get_column_letter(c)}{top + max(n, 1)}")
        for i in range(n):
            r = top + 1 + i
            for c, col in enumerate(cols, 1):
                cell = ws.cell(r, c)
                if col.get("formula"):
                    cell.value = col["formula"].format(r=r)
                elif i < len(data) and c - 1 < len(data[i]):
                    v = data[i][c - 1]
                    if col.get("type") == "date" and isinstance(v, str):
                        v = datetime.date.fromisoformat(v)
                    cell.value = v
                if col.get("format"):
                    cell.number_format = col["format"]
                elif col.get("type") == "date":
                    cell.number_format = "yyyy-mm-dd"
                if col.get("input"):
                    cell.fill = PatternFill("solid", fgColor=INPUT)
                cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if sh.get("totals") and n:
            r = top + 1 + n
            ws.cell(r, 1, "Total").font = Font(bold=True)
            for c, col in enumerate(cols, 1):
                fn = sh["totals"].get(col["header"])
                if fn:
                    L = get_column_letter(c)
                    cell = ws.cell(r, c, f"={fn}({L}{top + 1}:{L}{top + n})")
                    cell.font = Font(bold=True)
                    if col.get("format"):
                        cell.number_format = col["format"]
    if spec.get("sources"):
        ws = wb.create_sheet("Sources")
        ws.cell(1, 1, "Source").font = Font(bold=True, color=GOLD)
        ws.cell(1, 1).fill = PatternFill("solid", fgColor=BLACK)
        ws.column_dimensions["A"].width = 90
        for i, s in enumerate(spec["sources"], 2):
            ws.cell(i, 1, s)
    wb.save(out)


def main(a):
    if not a or a[0] in ("-h", "--help"):
        print(__doc__); return
    if a[0] == "--example":
        print(json.dumps(EXAMPLE, indent=1, ensure_ascii=False)); return
    spec = json.load(open(a[0], encoding="utf-8"))
    out = a[a.index("-o") + 1] if "-o" in a else os.path.splitext(a[0])[0] + ".xlsx"
    build(spec, out)
    print(f"wrote {out} ({os.path.getsize(out)} bytes) — sheets: {[s['name'] for s in spec['sheets']]}"
          f"{' + Sources' if spec.get('sources') else ''}")
    print(f"upload: create_file title={spec['title']!r} contentMimeType="
          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet base64Content=$(base64 -i <xlsx>) parentId=<folder>")


if __name__ == "__main__":
    main(sys.argv[1:])
